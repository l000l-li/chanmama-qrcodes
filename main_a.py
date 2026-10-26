# main_a.py
import logging
import os
import random
import re
import time

import requests
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

import config
from crawl import click_next_page, collect_product_hrefs, run_business_action
from login import create_logged_in_context
from parse import handle_detail_page
from utils import handle_slider_if_present, setup_logging

logger = logging.getLogger(__name__)

# 定义不抓取的黑名单品牌
BRAND_BLACKLIST = [
    "乔丹", "环球", "斐乐", "Fila", "足力健", "361°", "雅鹿", "特步",
    "回力", "西域骆驼", "安踏", "森马", "鸿星尔克", "李宁", "阿迪达斯",
    "adidas", "耐克", "NIKE", "KAPPA", "达芙妮", "奥古狮登"
]

EXCEL_PATH = config.EXCEL_PATH


def call_wps_action(action, **kwargs):
    """调用 AirScript；成功返回响应 dict，网络异常或云端报错返回 None。"""
    payload = {"Context": {"argv": {"action": action, **kwargs}}}
    headers = {"Content-Type": "application/json", "AirScript-Token": config.AIRSCRIPT_TOKEN}
    try:
        r = requests.post(config.AIRSCRIPT_WEBHOOK_A, json=payload, headers=headers, timeout=60)
        resp = r.json()
    except Exception as e:
        logger.error(f"网络请求异常: {repr(e)}")
        return None
    if resp.get("error"):
        logger.error(f"AirScript [{action}] 报错: {resp['error']}")
        return None
    return resp


def convert_single_item(raw_data):
    return [{
        "shopName": raw_data.get("店铺名称", ""),
        "imgUrl": raw_data.get("图片链接", ""),
        "link": raw_data.get("商品链接", ""),
        "listingDate": raw_data.get("上架时间", ""),
        "qrCodeUrl": raw_data.get("二维码URL", ""),
        "thirtyDaySales": raw_data.get("近30日销量", ""),
        "yesterdaySales": raw_data.get("昨日销量", "")
    }]


# ============ Excel 查重与记录 ============
def load_existing_links():
    if not os.path.exists(EXCEL_PATH):
        logger.warning(f"找不到本地记录表: {EXCEL_PATH}，将自动跳过查重。")
        return set()
    wb = load_workbook(EXCEL_PATH)
    ws = wb.active
    links = set()
    for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
        if row[0]:
            links.add(str(row[0]).strip())
    return links


def append_link_to_excel(link):
    """追加链接到本地查重表；成功返回 True，失败返回 False。"""
    if not os.path.exists(EXCEL_PATH):
        logger.error(f"找不到本地记录表: {EXCEL_PATH}，链接未记录: {link}")
        return False
    try:
        wb = load_workbook(EXCEL_PATH)
        ws = wb.active
        ws.append([link])
        wb.save(EXCEL_PATH)
        return True
    except Exception as e:
        logger.error(f"本地 Excel 写入失败（可能文件被占用），链接未记录，请手工补录: {link} | {repr(e)}")
        return False


# ============ 阶段 A 主流程 ============
def run_stage_a():
    logger.info("========== 阶段 A：发现新商品并追加 ==========")

    existing_links = load_existing_links()
    logger.info(f"✅ 已加载本地查重记录，共有 {len(existing_links)} 条历史链接。")

    with sync_playwright() as p:
        context = create_logged_in_context(p, config.RAW_COOKIE_PATH)
        page = context.new_page()

        logger.info("正在进入蝉妈妈首页...")
        page.goto(config.START_URL, timeout=30000)
        page.wait_for_load_state("networkidle")

        run_business_action(page)

        page_num = 1
        max_pages = 100
        stop_all = False

        detail = context.new_page()

        try:
            while page_num <= max_pages and not stop_all:
                logger.info(f"\n========== 第 {page_num} 页 ==========")
                page.wait_for_timeout(2000)
                hrefs = collect_product_hrefs(page)

                if not hrefs:
                    logger.info("本页无商品，结束")
                    break

                page_has_below_500 = False

                for idx, href in enumerate(hrefs, 1):
                    product_id = href.rstrip(".html").split("/")[-1]

                    try:
                        title = page.locator(f'a[href="{href}"]').filter(has_text=re.compile(r"\S")).first.inner_text()
                    except Exception:
                        title = "(未知)"

                    logger.info(f"  [{idx}/{len(hrefs)}] 检查: {title[:20]}...")

                    try:
                        # 1. 进详情页抓取
                        detail.goto(config.BASE_URL + href, timeout=30000)
                        detail.wait_for_load_state("domcontentloaded")
                        handle_slider_if_present(detail)

                        # 2. 获取数据
                        data = handle_detail_page(detail, product_id)
                        data["列表页标题"] = title

                        real_link = data.get("商品链接", "")
                        if not real_link:
                            logger.warning("     ⚠️ 未抓取到商品链接，跳过")
                            continue

                        clean_link = real_link.strip()

                        # ================= 品牌黑名单拦截 =================
                        shop_name = data.get("店铺名称", "").lower()  # 转小写，忽略大小写差异

                        # 检查店铺名是否包含黑名单里的任何一个词
                        hit_brand = next((brand for brand in BRAND_BLACKLIST if brand.lower() in shop_name), None)

                        if hit_brand:
                            logger.info(f"     🚫 触发品牌黑名单【{hit_brand}】，跳过店铺：{data.get('店铺名称')}")
                            continue  # 直接跳过，不写 WPS 也不写 Excel
                        # ===================================================

                        # 3. 检查是否在 Excel 里已存在
                        if clean_link in existing_links:
                            logger.info("     ⏭️ 本地 Excel 已有此链接，跳过")
                            continue

                        # 4. 检查销量
                        yesterday_text = data.get("昨日销量", "")
                        yesterday_num = None
                        m = re.match(r"([\d.]+)([wW万]?)", yesterday_text.lower().replace(",", ""))
                        if m:
                            num = float(m.group(1))
                            if m.group(2) in ("w", "万"):
                                num *= 10000
                            yesterday_num = num

                        if yesterday_num is not None and yesterday_num < 500:
                            page_has_below_500 = True

                        # 5. 写入 WPS（失败则不记为已采，下次继续尝试）
                        logger.info("     ✅ 发现新商品，正在写入 WPS...")
                        wps_data = convert_single_item(data)
                        if not call_wps_action("append_new_products", dataList=wps_data):
                            logger.warning("     ⚠️ WPS 写入失败，本条不计入本地记录，跳过")
                            continue

                        # 6. 写入本地 Excel 记录
                        if append_link_to_excel(clean_link):
                            existing_links.add(clean_link)
                            logger.info("     ✅ 已保存至本地记录表")
                        else:
                            logger.error(f"     ❌ WPS 已写入但本地记录失败，请手工补录: {clean_link}")

                    except Exception as e:
                        logger.warning(f"     ⚠️ 抓取失败: {repr(e)}")

                    time.sleep(random.uniform(3, 6))

                if page_has_below_500:
                    logger.info("✅ 本页出现昨日销量 <500 的商品，停止翻页。")
                    break

                if page_num >= max_pages:
                    break

                logger.info("准备翻下一页...")
                if not click_next_page(page):
                    logger.info("没有下一页了，结束。")
                    break

                page_num += 1
                time.sleep(random.uniform(3, 6))

        finally:
            detail.close()
            context.browser.close()

    logger.info("✅ 阶段 A 完成！")


if __name__ == "__main__":
    setup_logging()
    run_stage_a()
